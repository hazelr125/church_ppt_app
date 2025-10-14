#app.py
import os
from flask import Flask, request, send_file, render_template, flash, redirect, url_for
from build_helpers import parse_pdf_to_structured, build_mapping_wrapper
from generate_ppt import generate_presentation
from hymns_db import HymnDatabase, process_user_hymns
from docx import Document
import re
from parse_announcements import parse_announcements_docx

UPLOAD_DIR = os.path.join(os.getcwd(), 'uploads')
os.makedirs(UPLOAD_DIR, exist_ok=True)

app = Flask(__name__)
app.secret_key = "change-this-secret-for-production"

hymn_db = HymnDatabase("kannada_hymns.csv", "tulu_hymns.csv", "english_hymns.csv")

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        pdf_file = request.files.get('pdf')
        tpl_file = request.files.get('tpl')
        ann_doc_file = request.files.get('ann_doc')

        if not pdf_file or not tpl_file:
            flash("Please upload both the weekly PDF and a Canva PPTX template.")
            return redirect(url_for('index'))
        
        hymn_placeholders = process_user_hymns(request.form, hymn_db)

        print(f"Generated {len(hymn_placeholders)} hymn placeholders")
        for key, value in list(hymn_placeholders.items())[:5]:  # Show first 5
            print(f"  {key}: {str(value)[:50]}...")

        pdf_path = os.path.join(UPLOAD_DIR, 'input.pdf'); pdf_file.save(pdf_path)
        tpl_path = os.path.join(UPLOAD_DIR, 'template.pptx'); tpl_file.save(tpl_path)

        parsed = parse_pdf_to_structured(pdf_path)

        mapping, announcements_table = build_mapping_wrapper(parsed, hymn_db)
        mapping.update({f'{{{k}}}': v for k, v in hymn_placeholders.items()})
        
        if ann_doc_file:
            ann_path = os.path.join(UPLOAD_DIR, 'announcements.docx')
            ann_doc_file.save(ann_path)
            ann_table, ann_text = parse_announcements_docx(ann_path)

    # Override announcements text
            parsed['announcements_block'] = ann_text
            mapping['{ANNOUNCEMENTS_TEXT}'] = ann_text

    # Override announcements table if we found one
        if ann_table:
            announcements_table = ann_table

        out_pptx = os.path.join(UPLOAD_DIR, 'final_presentation.pptx')
        generate_presentation(tpl_path, out_pptx, mapping, announcements_table_data=announcements_table)

        return send_file(out_pptx, as_attachment=True, download_name='final_presentation.pptx')

    return render_template('index.html')

if __name__ == "__main__":
    print("Starting Church PPT Builder on http://127.0.0.1:5000")
    app.run(debug=True, host="127.0.0.1", port=5000)
